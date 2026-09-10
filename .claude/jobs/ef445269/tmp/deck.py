from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

INK=RGBColor(0x1B,0x26,0x35); BODY=RGBColor(0x3F,0x4C,0x5C)
ACCENT=RGBColor(0x0F,0x9E,0x8E); MUTED=RGBColor(0x8A,0x97,0xA6); WHITE=RGBColor(0xFF,0xFF,0xFF)
LIGHT=RGBColor(0xB8,0xC4,0xD0)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
def slide(): return prs.slides.add_slide(BLANK)
def bg(s,c):
    s.background.fill.solid(); s.background.fill.fore_color.rgb=c
def text(s,txt,l,t,w,h,size,c=INK,bold=False,align=PP_ALIGN.LEFT,italic=False,sp=1.0):
    tf=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)).text_frame
    tf.word_wrap=True
    for i,line in enumerate(txt.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.alignment=align; p.line_spacing=sp
        for r in p.runs:
            r.font.size=Pt(size); r.font.color.rgb=c; r.font.bold=bold
            r.font.italic=italic; r.font.name="Segoe UI"
def rule(s,top,left=0.9,width=2.2,c=ACCENT):
    b=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(left),Inches(top),Inches(width),Pt(3))
    b.fill.solid(); b.fill.fore_color.rgb=c; b.line.fill.background(); b.shadow.inherit=False
def kicker(s,t): text(s,t.upper(),0.9,0.55,11.5,0.4,13,ACCENT,bold=True)
def section(s,num,title,sub=""):
    bg(s,INK); text(s,num,0.9,2.1,3,1.0,15,ACCENT,bold=True)
    text(s,title,0.9,2.6,11.5,1.6,46,WHITE,bold=True)
    if sub: text(s,sub,0.9,4.3,10.5,1.2,19,LIGHT)
def split(s,word,no,yes):
    kicker(s,"what i mean by"); text(s,word,0.9,1.05,11.5,1.0,40,INK,bold=True); rule(s,2.15)
    text(s,"I don't mean",0.9,2.6,5.2,0.4,14,MUTED,bold=True)
    text(s,no,0.9,3.1,5.2,2.4,21,BODY,sp=1.2)
    text(s,"I do mean",6.9,2.6,5.5,0.4,14,ACCENT,bold=True)
    text(s,yes,6.9,3.1,5.5,2.4,21,INK,bold=True,sp=1.2)

s=slide(); bg(s,INK)
text(s,"WAYS OF WORKING",0.9,2.3,11.5,0.5,15,ACCENT,bold=True)
text(s,"How we work together,\nand where things live",0.9,2.85,11.5,2.0,44,WHITE,bold=True,sp=1.15)
text(s,"Digital team  ·  10 September 2026",0.9,5.5,11.5,0.5,17,LIGHT)

s=slide(); kicker(s,"why we're here")
text(s,"“You and I are drowning in no structure.\nAnd if we are drowning, what will\nour team feel like?”",0.9,1.9,11.5,2.6,33,INK,sp=1.2)
rule(s,4.85); text(s,"Erica, earlier this week",0.9,5.1,6,0.4,17,MUTED)

s=slide(); kicker(s,"first, what this is not")
text(s,"Not restructuring 776 work items\n\nNot changing tooling today\n\nNot asking anyone to redo past work",0.9,1.5,11.5,2.8,26,BODY,sp=1.15)
rule(s,4.7)
text(s,"What it is: four things agreed, so everyone knows where work goes\nand what happens when.",0.9,5.0,11.0,1.2,24,INK,bold=True,sp=1.2)

s=slide(); kicker(s,"before we start")
text(s,"Invisible drag",0.9,1.05,11.5,0.9,40,INK,bold=True); rule(s,2.05)
text(s,"Everyone nods. Everyone agrees.\nThe work goes in different directions anyway.",0.9,2.5,11.0,1.4,27,BODY,sp=1.2)
text(s,"Usually because one word meant different things to different people.\nSo here are the five that matter today.",0.9,4.4,11.0,1.4,21,MUTED,sp=1.25)

split(slide(),"Structure","Process for its own sake.\nMore approvals to get through.","Everyone knowing where a thing\ngoes, and who picks it up.")
split(slide(),"Ceremony","More meetings.","Fewer meetings, with a written\npurpose, that end in something\nchanging on the board.")
split(slide(),"Feature","Everything you work on.","The unit we prioritise and track\nvalue against. Your day-to-day\nsits underneath it, and that's fine.")
split(slide(),"Ways of working","A standard you get measured\nagainst.","The small number of agreements\nthat stop us tripping over\neach other.")
split(slide(),"The board","It could mean four things —\nthe tool, what's typed into it,\nwhether people update it…","…or whether it reflects what's\nactually true. I mean that one.\nOthers may mean a different one.")

section(slide(),"01","Teams structure","Where things live, and which door to knock on  ·  15 min")
s=slide(); kicker(s,"01 · teams structure")
text(s,"Three departments",0.9,1.0,11.5,0.8,34,INK,bold=True); rule(s,1.9)
y=2.35
for n,d in [("Inform Digital","Internal IT and management escalation"),
            ("Digital Capability","Everything that supports delivery to clients,\nbut not client delivery"),
            ("Client Delivery","The engagements. Already exists per client")]:
    text(s,n,0.9,y,3.6,0.5,22,ACCENT,bold=True); text(s,d,4.7,y,7.6,1.0,19,BODY,sp=1.15); y+=1.45
text(s,"Two shop windows, because the audiences differ.",0.9,6.4,11,0.5,20,INK,bold=True)

section(slide(),"02","Ceremonies","Why, when, how often, and what comes out  ·  15 min")
s=slide(); kicker(s,"02 · ceremonies")
text(s,"A ceremony that doesn't change\nthe board is a conversation.",0.9,1.0,11.5,1.6,32,INK,bold=True,sp=1.15)
rule(s,2.9); y=3.35
for n,f,w in [("Engineering check-in","Daily, short","Unblocks fast"),
              ("Delivery review","Weekly","Features move, priorities set"),
              ("Ways of working","Monthly","Where this gets amended")]:
    text(s,n,0.9,y,4.2,0.5,21,ACCENT,bold=True); text(s,f,5.3,y,2.2,0.5,19,INK,bold=True)
    text(s,w,7.7,y,4.6,0.5,19,BODY); y+=1.0
text(s,"Every one gets a written purpose and an agenda.",0.9,6.5,11,0.5,20,MUTED,italic=True)

section(slide(),"03","Backlogs and DevOps","How work is tracked, and how it fits together  ·  15 min")
s=slide(); kicker(s,"03 · the four levels")
text(s,"Product Family  ›  Product  ›  Feature  ›  PBI",0.9,1.3,11.8,0.9,27,INK,bold=True); rule(s,2.35)
text(s,"Features are what we prioritise and track value against.\n\nPBIs are how the team gets a feature done — that's the team\ndoing their work, and it doesn't need steering from above.",0.9,2.85,11.2,2.4,23,BODY,sp=1.25)

s=slide(); kicker(s,"03 · what changes on the board")
text(s,"Six changes. Not 776.",0.9,1.0,11.5,0.8,34,INK,bold=True); rule(s,1.9)
text(s,"Three epics for the three blocks\n\nResolve the AI strategy duplicates\n\nOwn the capability build\n\nAdd the missing epics\n\nCollapse legacy behind a label\n\nDecide if Frontier Firm gets an item",0.9,2.4,11.0,4.2,21,BODY,sp=1.1)

section(slide(),"04","Ways of working","The small number of agreements  ·  15 min")
s=slide(); kicker(s,"04 · how we'll go about it")
text(s,"“Keep it simple to start with and\nbuild on it. Build the basics,\niterate from there.”",0.9,1.4,11.5,2.2,32,INK,sp=1.2)
rule(s,4.1); text(s,"Erica",0.9,4.35,6,0.4,17,MUTED)
text(s,"On tooling: this is not a decision. It's a decision to stop arguing\nabout tools until we have numbers. DevOps is the floor, not the winner.",0.9,5.25,11.2,1.3,20,BODY,sp=1.2)

s=slide(); bg(s,INK)
text(s,"WHAT CHANGES TOMORROW",0.9,0.9,11.5,0.5,15,ACCENT,bold=True)
text(s,"Every recurring meeting gets a purpose and an agenda\n\nFeatures are what we prioritise; PBIs are how work gets done\n\nIf we discuss it, it lands on the board\n\nWhen we find a gap, the first question is “was this ever\nanyone's job” — not “who dropped this”\n\nNothing else changes this week",0.9,1.8,11.5,4.6,23,WHITE,sp=1.15)

s=slide(); kicker(s,"over to you")
text(s,"What's missing,\nand what's too much?",0.9,2.2,11.5,2.0,42,INK,bold=True,sp=1.15)
rule(s,4.6)
text(s,"Nothing here is finished. If there's a better way to get there,\nI'd rather have that than my version.",0.9,5.0,11.2,1.3,22,BODY,sp=1.2)

out="/home/jellypai/corpus/inform/Working_area/ways_of_working_2026-09-10.pptx"
prs.save(out); print("saved:",out); print("slides:",len(prs.slides._sldIdLst))
